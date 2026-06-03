"""
Unit tests for document parsers.
"""
import re
import tempfile
from pathlib import Path

import pytest
from docx import Document

from parsers.docx_parser import DocxParser, _detect_section, _is_heading, _normalize
from parsers.text_parser import (
    TextParser,
    _extract_job_title,
    _extract_keywords,
    _extract_required_skills,
    _extract_responsibilities,
    _extract_skills_from_text,
)
from tests.conftest import SAMPLE_JD_TEXT, SAMPLE_RESUME_TEXT


# ---------------------------------------------------------------------------
# DocxParser helpers
# ---------------------------------------------------------------------------

class TestDocxParserHelpers:
    def test_normalize(self):
        assert _normalize("  EXPERIENCE  ") == "experience"
        assert _normalize("Work History") == "work history"

    def test_detect_section_known(self):
        assert _detect_section("Work Experience") == "experience"
        assert _detect_section("Education") == "education"
        assert _detect_section("Technical Skills") == "skills"
        assert _detect_section("Professional Summary") == "summary"
        assert _detect_section("Certifications") == "certifications"
        assert _detect_section("Projects") == "projects"

    def test_detect_section_unknown(self):
        assert _detect_section("Random Header XYZ") is None
        assert _detect_section("") is None

    def test_is_heading_with_style(self, tmp_path):
        doc = Document()
        para = doc.add_heading("Experience", level=2)
        assert _is_heading(para) is True

    def test_is_heading_bold_short(self, tmp_path):
        doc = Document()
        para = doc.add_paragraph()
        run = para.add_run("SKILLS")
        run.bold = True
        assert _is_heading(para) is True

    def test_is_heading_normal_long_text(self, tmp_path):
        doc = Document()
        para = doc.add_paragraph(
            "This is a normal paragraph that is definitely not a heading."
        )
        assert _is_heading(para) is False


class TestDocxParser:
    def test_parse_returns_correct_type(self, tmp_docx_resume):
        from core.models import DocumentType
        result = DocxParser().parse(tmp_docx_resume)
        assert result.source_type == DocumentType.DOCX

    def test_parse_extracts_sections(self, tmp_docx_resume):
        result = DocxParser().parse(tmp_docx_resume)
        section_names = [s.name for s in result.sections]
        assert "experience" in section_names
        assert "skills" in section_names
        assert "education" in section_names

    def test_parse_extracts_contact_email(self, tmp_docx_resume):
        result = DocxParser().parse(tmp_docx_resume)
        assert "email" in result.contact_info
        assert "@" in result.contact_info["email"]

    def test_parse_extracts_contact_name(self, tmp_docx_resume):
        result = DocxParser().parse(tmp_docx_resume)
        # Name should be extracted from top lines
        assert "name" in result.contact_info
        assert result.contact_info["name"] == "Jane Doe"

    def test_parse_style_metadata(self, tmp_docx_resume):
        result = DocxParser().parse(tmp_docx_resume)
        assert result.style_metadata is not None
        assert isinstance(result.style_metadata.fonts, list)

    def test_parse_full_text_nonempty(self, tmp_docx_resume):
        result = DocxParser().parse(tmp_docx_resume)
        assert len(result.full_text) > 50

    def test_extract_text_method(self, tmp_docx_resume):
        text = DocxParser().extract_text(tmp_docx_resume)
        assert "Jane Doe" in text or "Engineer" in text

    def test_sections_have_raw_text(self, tmp_docx_resume):
        result = DocxParser().parse(tmp_docx_resume)
        for section in result.sections:
            assert isinstance(section.raw_text, str)

    def test_model_copy_deep(self, tmp_docx_resume):
        """Ensure deep copy doesn't mutate original (critical for pipeline)."""
        result = DocxParser().parse(tmp_docx_resume)
        copy = result.model_copy(deep=True)
        copy.sections[0].rewritten_text = "MODIFIED"
        assert result.sections[0].rewritten_text is None


# ---------------------------------------------------------------------------
# TextParser
# ---------------------------------------------------------------------------

class TestTextParserHelpers:
    def test_extract_job_title(self):
        title = _extract_job_title(SAMPLE_JD_TEXT)
        assert "Senior" in title or "Engineer" in title

    def test_extract_skills_from_text(self):
        skills = _extract_skills_from_text("We need python and kubernetes experts")
        assert "python" in skills
        assert "kubernetes" in skills

    def test_extract_required_skills(self):
        skills = _extract_required_skills(SAMPLE_JD_TEXT)
        assert len(skills) > 0
        # Should find bullet-pointed items in Requirements section
        all_skills = " ".join(skills).lower()
        assert "python" in all_skills or "kubernetes" in all_skills

    def test_extract_responsibilities(self):
        resp = _extract_responsibilities(SAMPLE_JD_TEXT)
        assert len(resp) > 0
        assert any("pipeline" in r.lower() or "AI" in r or "LLM" in r for r in resp)

    def test_extract_keywords(self):
        keywords = _extract_keywords(SAMPLE_JD_TEXT)
        assert len(keywords) > 0


class TestTextParser:
    def test_parse_returns_parsed_jd(self):
        from core.models import ParsedJD
        result = TextParser().parse(SAMPLE_JD_TEXT)
        assert isinstance(result, ParsedJD)

    def test_parse_extracts_title(self):
        result = TextParser().parse(SAMPLE_JD_TEXT)
        assert len(result.job_title) > 0

    def test_parse_extracts_required_skills(self):
        result = TextParser().parse(SAMPLE_JD_TEXT)
        assert len(result.required_skills) > 0

    def test_parse_extracts_responsibilities(self):
        result = TextParser().parse(SAMPLE_JD_TEXT)
        assert len(result.responsibilities) > 0

    def test_parse_extracts_keywords(self):
        result = TextParser().parse(SAMPLE_JD_TEXT)
        assert len(result.keywords) > 0

    def test_parse_extracts_preferred_skills(self):
        result = TextParser().parse(SAMPLE_JD_TEXT)
        assert len(result.preferred_skills) > 0

    def test_parse_empty_text(self):
        result = TextParser().parse("")
        assert result.full_text == ""
        assert result.required_skills == []

    def test_parse_preserves_full_text(self):
        result = TextParser().parse(SAMPLE_JD_TEXT)
        assert result.full_text == SAMPLE_JD_TEXT
