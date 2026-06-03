"""
API integration tests using FastAPI TestClient.
The LLM rewriter and embedding model are mocked for speed.
"""
import io
import json
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

import pytest
from docx import Document
from fastapi.testclient import TestClient

from tests.conftest import SAMPLE_JD_TEXT


# ---------------------------------------------------------------------------
# App fixture with mocked expensive components
# ---------------------------------------------------------------------------

@pytest.fixture(scope="module")
def client(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("api_test")

    # Patch settings dirs to use tmp
    with patch("core.config.Settings.ensure_dirs"):
        with patch("core.config.get_settings") as mock_settings:
            from core.config import Settings
            s = Settings(
                upload_dir=tmp / "uploads",
                output_dir=tmp / "outputs",
                ollama_model="llama3.1:8b",
            )
            s.upload_dir.mkdir(parents=True, exist_ok=True)
            s.output_dir.mkdir(parents=True, exist_ok=True)
            mock_settings.return_value = s

            # Patch embedding model loading so tests don't download
            with patch("nlp.embedder.get_embedder") as mock_emb:
                import numpy as np
                mock_model = MagicMock()
                mock_model.encode.return_value = np.random.rand(1, 384).astype("float32")
                mock_emb.return_value = mock_model

                from main import app
                yield TestClient(app, raise_server_exceptions=False)


@pytest.fixture
def docx_resume_bytes(tmp_path):
    doc = Document()
    doc.add_heading("Jane Doe", level=1)
    doc.add_paragraph("jane@example.com | +91-9876543210")
    doc.add_heading("Experience", level=2)
    doc.add_paragraph("Senior Software Engineer at Acme Corp")
    doc.add_heading("Skills", level=2)
    doc.add_paragraph("Python, Kubernetes, FastAPI, Docker")
    path = tmp_path / "resume.docx"
    doc.save(str(path))
    return path.read_bytes()


# ---------------------------------------------------------------------------
# Health endpoint
# ---------------------------------------------------------------------------

class TestHealthEndpoint:
    def test_health_returns_200(self, client):
        with patch("routes.health.httpx.AsyncClient") as mock_client_class:
            mock_client = AsyncMock()
            mock_resp = MagicMock()
            mock_resp.status_code = 200
            mock_resp.json.return_value = {"models": [{"name": "llama3.1:8b"}]}
            mock_client.__aenter__ = AsyncMock(return_value=mock_client)
            mock_client.__aexit__ = AsyncMock(return_value=None)
            mock_client.get = AsyncMock(return_value=mock_resp)
            mock_client_class.return_value = mock_client

            resp = client.get("/api/v1/health")

        assert resp.status_code == 200
        data = resp.json()
        assert "status" in data
        assert "version" in data

    def test_health_degraded_when_ollama_down(self, client):
        with patch("routes.health.httpx.AsyncClient") as mock_client_class:
            mock_client = AsyncMock()
            mock_client.__aenter__ = AsyncMock(return_value=mock_client)
            mock_client.__aexit__ = AsyncMock(return_value=None)
            mock_client.get = AsyncMock(side_effect=Exception("Connection refused"))
            mock_client_class.return_value = mock_client

            resp = client.get("/api/v1/health")

        assert resp.status_code == 200
        data = resp.json()
        assert data["ollama_connected"] is False


# ---------------------------------------------------------------------------
# Tailor endpoint
# ---------------------------------------------------------------------------

class TestTailorEndpoint:
    def test_tailor_requires_resume(self, client):
        resp = client.post("/api/v1/tailor", data={"jd_text": "Some JD"})
        assert resp.status_code == 422  # FastAPI validation

    def test_tailor_requires_jd(self, client, docx_resume_bytes):
        resp = client.post(
            "/api/v1/tailor",
            files={"resume": ("resume.docx", docx_resume_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        assert resp.status_code == 400
        assert "JD" in resp.json().get("detail", "")

    def test_tailor_accepts_docx_with_jd_text(self, client, docx_resume_bytes):
        resp = client.post(
            "/api/v1/tailor",
            files={"resume": ("resume.docx", docx_resume_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data={"jd_text": SAMPLE_JD_TEXT},
        )
        assert resp.status_code == 200
        data = resp.json()
        assert "job_id" in data
        assert data["status"] == "queued"

    def test_tailor_rejects_invalid_extension(self, client):
        resp = client.post(
            "/api/v1/tailor",
            files={"resume": ("resume.txt", b"resume content", "text/plain")},
            data={"jd_text": "some jd"},
        )
        assert resp.status_code == 400

    def test_tailor_returns_unique_job_ids(self, client, docx_resume_bytes):
        def submit():
            return client.post(
                "/api/v1/tailor",
                files={"resume": ("resume.docx", docx_resume_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                data={"jd_text": SAMPLE_JD_TEXT},
            ).json()["job_id"]

        j1, j2 = submit(), submit()
        assert j1 != j2


# ---------------------------------------------------------------------------
# Jobs endpoints
# ---------------------------------------------------------------------------

class TestJobsEndpoints:
    def _create_job(self, client, docx_resume_bytes) -> str:
        resp = client.post(
            "/api/v1/tailor",
            files={"resume": ("resume.docx", docx_resume_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data={"jd_text": SAMPLE_JD_TEXT},
        )
        return resp.json()["job_id"]

    def test_result_returns_status_for_queued_job(self, client, docx_resume_bytes):
        job_id = self._create_job(client, docx_resume_bytes)
        resp = client.get(f"/api/v1/jobs/{job_id}/result")
        assert resp.status_code == 200
        data = resp.json()
        assert data["job_id"] == job_id
        assert "status" in data

    def test_result_404_for_unknown_job(self, client):
        resp = client.get("/api/v1/jobs/nonexistent-job-id/result")
        assert resp.status_code == 404

    def test_download_endpoints_exist(self, client, docx_resume_bytes):
        """Download endpoints should return files or meaningful errors, never 500."""
        job_id = self._create_job(client, docx_resume_bytes)
        # Either complete (200) or not-ready (400/404) — never crash
        for endpoint in ["docx", "pdf"]:
            resp = client.get(f"/api/v1/jobs/{job_id}/download/{endpoint}")
            assert resp.status_code in (200, 400, 404)

    def test_delete_job(self, client, docx_resume_bytes):
        job_id = self._create_job(client, docx_resume_bytes)
        resp = client.delete(f"/api/v1/jobs/{job_id}")
        assert resp.status_code == 200
        assert resp.json()["deleted"] is True

    def test_delete_nonexistent_job_404(self, client):
        resp = client.delete("/api/v1/jobs/ghost-job")
        assert resp.status_code == 404

    def test_root_endpoint(self, client):
        resp = client.get("/")
        assert resp.status_code == 200
        assert "careerfit-ai" in resp.json().get("service", "")


# ---------------------------------------------------------------------------
# Pipeline integration (mocked LLM + embedder)
# ---------------------------------------------------------------------------

class TestPipelineIntegration:
    @pytest.mark.asyncio
    async def test_full_pipeline_with_mocks(self, tmp_path):
        """Run the full pipeline end-to-end with mocked LLM and embeddings."""
        import numpy as np
        from core.job_manager import JobManager
        from core.pipeline import run_pipeline
        from docx import Document

        # Create a real DOCX resume
        doc = Document()
        doc.add_heading("Jane Doe", level=1)
        doc.add_paragraph("jane@test.com | +91-9876543210")
        doc.add_heading("Experience", level=2)
        doc.add_paragraph("Senior Software Engineer — Built Python APIs and Kubernetes deployments")
        doc.add_heading("Skills", level=2)
        doc.add_paragraph("Python, Kubernetes, Docker, FastAPI")
        resume_path = tmp_path / "resume.docx"
        doc.save(str(resume_path))

        # Set up settings with tmp dirs
        with patch("core.config.get_settings") as mock_gs:
            from core.config import Settings
            s = Settings(upload_dir=tmp_path / "up", output_dir=tmp_path / "out")
            s.upload_dir.mkdir(parents=True, exist_ok=True)
            s.output_dir.mkdir(parents=True, exist_ok=True)
            mock_gs.return_value = s

            jm = JobManager()
            job = jm.create_job(resume_path=resume_path, jd_text=SAMPLE_JD_TEXT)

            # Mock embedder
            with patch("nlp.embedder.embed") as mock_embed:
                mock_embed.return_value = np.random.rand(1, 384).astype("float32")

                # Mock LLM rewriter
                with patch("generators.llm_rewriter.LLMRewriter.rewrite_section",
                           new_callable=AsyncMock) as mock_rw:
                    mock_rw.return_value = "Improved rewritten text with Kubernetes and FastAPI."

                    await run_pipeline(job)

        from core.models import JobStatus
        assert job.status == JobStatus.COMPLETE
        assert job.result is not None
        assert job.result.docx_url is not None
        assert job.result.pdf_url is not None
        assert job.result.gap_report is not None
