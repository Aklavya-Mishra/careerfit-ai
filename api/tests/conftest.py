"""
Shared fixtures for all test modules.
"""
import io
import tempfile
from pathlib import Path

import pytest
from docx import Document

# ---------------------------------------------------------------------------
# Sample data constants
# ---------------------------------------------------------------------------

SAMPLE_RESUME_TEXT = """
Jane Doe
jane.doe@email.com | +91-9876543210 | linkedin.com/in/janedoe

SUMMARY
Experienced software engineer with 6 years building distributed systems and cloud-native applications.
Proficient in Python, Go, and Kubernetes. Strong background in API design and microservices.

EXPERIENCE
Senior Software Engineer — Acme Corp (2021–Present)
- Led migration of monolithic app to microservices architecture, reducing latency by 40%
- Built real-time data pipelines processing 100k events/sec using Kafka and Python
- Mentored a team of 4 junior engineers

Software Engineer — StartupXYZ (2019–2021)
- Developed RESTful APIs using FastAPI and PostgreSQL
- Containerized applications using Docker and deployed to GCP

EDUCATION
B.E. Computer Science — BITS Pilani (2015–2019)

SKILLS
Python, Go, Kubernetes, Docker, Kafka, FastAPI, PostgreSQL, Redis, GCP, Git, CI/CD
""".strip()

SAMPLE_JD_TEXT = """
Senior AI Systems Engineer

We are looking for a Senior AI Systems Engineer with 5+ years of experience.

Requirements:
• 5+ years of Python experience
• Strong knowledge of LLM orchestration and AI pipelines
• Experience with Kubernetes, Docker, and cloud platforms (GCP/Azure)
• Proficiency in FastAPI or similar frameworks
• Experience with vector databases (Pinecone, Weaviate)
• Knowledge of MLOps practices

Responsibilities:
• Design and implement real-time AI inference pipelines
• Integrate and manage LLM APIs (OpenAI, Anthropic, local models)
• Build robust monitoring and observability for AI systems
• Collaborate with ML engineers and product teams

Preferred:
• Experience with speech-to-text systems
• Knowledge of RAG architectures
• Familiarity with Apache Kafka for streaming
""".strip()


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture
def sample_resume_text():
    return SAMPLE_RESUME_TEXT


@pytest.fixture
def sample_jd_text():
    return SAMPLE_JD_TEXT


@pytest.fixture
def tmp_docx_resume(tmp_path):
    """Create a real DOCX file with resume content."""
    doc = Document()
    doc.add_heading("Jane Doe", level=1)
    doc.add_paragraph("jane.doe@email.com | +91-9876543210")
    doc.add_heading("Summary", level=2)
    doc.add_paragraph(
        "Experienced software engineer with 6 years building distributed systems."
    )
    doc.add_heading("Experience", level=2)
    doc.add_paragraph("Senior Software Engineer — Acme Corp (2021–Present)")
    p = doc.add_paragraph()
    p.add_run("Led migration to microservices, reducing latency by 40%")
    doc.add_heading("Skills", level=2)
    doc.add_paragraph("Python, Kubernetes, Docker, FastAPI, PostgreSQL, GCP")
    doc.add_heading("Education", level=2)
    doc.add_paragraph("B.E. Computer Science — BITS Pilani (2019)")

    path = tmp_path / "resume.docx"
    doc.save(str(path))
    return path


@pytest.fixture
def tmp_text_jd(tmp_path):
    """Write JD to a temp text file."""
    path = tmp_path / "jd.txt"
    path.write_text(SAMPLE_JD_TEXT)
    return path
