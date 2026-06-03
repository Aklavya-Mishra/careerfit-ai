"""
Unit tests for document generators (DOCX and PDF).
Ollama LLM rewriter is mocked to avoid requiring a live Ollama instance.
"""
import tempfile
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

import pytest
from docx import Document

from core.models import (
    DocumentType,
    GapReport,
    ParsedJD,
    ParsedResume,
    ResumeSection,
    SectionGapScore,
    StyleMetadata,
)
from generators.docx_generator import DocxGenerator
from generators.pdf_generator import PdfGenerator, _build_html, _pick_accent_color, _pick_font
from tests.conftest import SAMPLE_JD_TEXT


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def make_tailored_resume(with_rewrites=True) -> ParsedResume:
    sections = [
        ResumeSection(
            name="summary",
            raw_text="Software engineer with 6 years experience in Python and distributed systems.",
            rewritten_text=(
                "Senior AI Systems Engineer with 6 years building scalable ML pipelines "
                "and LLM orchestration systems using Python, Kubernetes, and FastAPI."
                if with_rewrites else None
            ),
        ),
        ResumeSection(
            name="experience",
            raw_text="Led migration of monolithic app to microservices.",
            rewritten_text=(
                "Led migration to microservices architecture, improving real-time AI inference "
                "latency by 40% on GCP using Kubernetes and FastAPI."
                if with_rewrites else None
            ),
        ),
        ResumeSection(
            name="skills",
            raw_text="Python, Docker, Kubernetes, FastAPI, PostgreSQL",
            rewritten_text=(
                "Python, Kubernetes, Docker, FastAPI, LLM orchestration, GCP, MLOps, PostgreSQL, Redis"
                if with_rewrites else None
            ),
        ),
    ]
    return ParsedResume(
        source_type=DocumentType.DOCX,
        full_text="Jane Doe resume text...",
        sections=sections,
        style_metadata=StyleMetadata(
            fonts=["Calibri"],
            colors=["#1F579C"],
            has_tables=False,
        ),
        contact_info={
            "name": "Jane Doe",
            "email": "jane.doe@email.com",
            "phone": "+91-9876543210",
        },
    )


# ---------------------------------------------------------------------------
# PDF Generator helpers
# ---------------------------------------------------------------------------

class TestPdfGeneratorHelpers:
    def test_pick_font_calibri(self):
        style = StyleMetadata(fonts=["CalibriRegular", "CalibriLight"])
        assert "Calibri" in _pick_font(style)

    def test_pick_font_fallback(self):
        style = StyleMetadata(fonts=["UnknownFont123"])
        font = _pick_font(style)
        assert font == "Arial"

    def test_pick_font_empty(self):
        style = StyleMetadata(fonts=[])
        assert _pick_font(style) == "Arial"

    def test_pick_accent_color_valid(self):
        style = StyleMetadata(colors=["#2B6CB0", "#FFFFFF"])
        color = _pick_accent_color(style)
        assert color == "#2B6CB0"

    def test_pick_accent_color_skips_black_white(self):
        style = StyleMetadata(colors=["#000000", "#FFFFFF"])
        color = _pick_accent_color(style)
        assert color == "#1F579C"  # fallback blue

    def test_pick_accent_color_fallback(self):
        style = StyleMetadata(colors=[])
        assert _pick_accent_color(style) == "#1F579C"

    def test_build_html_contains_name(self):
        resume = make_tailored_resume()
        html = _build_html(resume)
        assert "Jane Doe" in html

    def test_build_html_contains_section_content(self):
        resume = make_tailored_resume()
        html = _build_html(resume)
        assert "Kubernetes" in html or "Python" in html

    def test_build_html_contains_font(self):
        resume = make_tailored_resume()
        html = _build_html(resume)
        assert "Calibri" in html

    def test_build_html_contains_accent_color(self):
        resume = make_tailored_resume()
        html = _build_html(resume)
        assert "#1F579C" in html or "#2B6CB0" in html or "#1f579c" in html.lower()

    def test_build_html_is_valid_html(self):
        resume = make_tailored_resume()
        html = _build_html(resume)
        assert html.startswith("<!DOCTYPE html>")
        assert "</html>" in html

    def test_build_html_uses_rewritten_text(self):
        resume = make_tailored_resume(with_rewrites=True)
        html = _build_html(resume)
        assert "LLM orchestration" in html

    def test_build_html_falls_back_to_raw_text(self):
        resume = make_tailored_resume(with_rewrites=False)
        html = _build_html(resume)
        assert "Python, Docker" in html or "microservices" in html


# ---------------------------------------------------------------------------
# PDF Generator file output
# ---------------------------------------------------------------------------

class TestPdfGenerator:
    def test_generate_creates_file(self, tmp_path):
        resume = make_tailored_resume()
        gen = PdfGenerator()
        out = gen.generate(tailored_resume=resume, output_dir=tmp_path, job_id="test123")
        assert out.exists()
        assert out.suffix == ".pdf"
        assert out.stat().st_size > 100

    def test_generate_correct_filename(self, tmp_path):
        resume = make_tailored_resume()
        gen = PdfGenerator()
        out = gen.generate(tailored_resume=resume, output_dir=tmp_path, job_id="abc456")
        assert "abc456" in out.name
        assert "tailored" in out.name


# ---------------------------------------------------------------------------
# DOCX Generator
# ---------------------------------------------------------------------------

class TestDocxGenerator:
    def test_generate_new_creates_file(self, tmp_path):
        resume = make_tailored_resume()
        gen = DocxGenerator()
        out = gen.generate(
            tailored_resume=resume,
            original_path=None,
            output_dir=tmp_path,
            job_id="test789",
        )
        assert out.exists()
        assert out.suffix == ".docx"
        assert out.stat().st_size > 500

    def test_generate_new_docx_readable(self, tmp_path):
        resume = make_tailored_resume()
        gen = DocxGenerator()
        out = gen.generate(
            tailored_resume=resume,
            original_path=None,
            output_dir=tmp_path,
            job_id="readtest",
        )
        doc = Document(str(out))
        full = "\n".join(p.text for p in doc.paragraphs)
        assert len(full) > 50

    def test_generate_from_original_preserves_structure(self, tmp_path, tmp_docx_resume):
        resume = make_tailored_resume()
        # Match section names to original doc
        gen = DocxGenerator()
        out = gen.generate(
            tailored_resume=resume,
            original_path=tmp_docx_resume,
            output_dir=tmp_path,
            job_id="preserve_test",
        )
        assert out.exists()
        assert out.suffix == ".docx"
        # Should be based on original
        doc = Document(str(out))
        assert len(doc.paragraphs) > 0

    def test_generate_from_original_preserves_unmatched_bullets(self, tmp_path):
        original_second_bullet = (
            "Troubleshot and resolved customer-impacting issues across orders, "
            "deliveries, returns, and account workflows."
        )
        doc = Document()
        doc.add_heading("Experience", level=2)
        doc.add_paragraph("Customer Support Associate")
        doc.add_paragraph(
            "Assisted customers by providing timely support through phone, email, "
            "and chat channels."
        )
        doc.add_paragraph(original_second_bullet)
        original_path = tmp_path / "original.docx"
        doc.save(str(original_path))

        resume = ParsedResume(
            source_type=DocumentType.DOCX,
            full_text="",
            sections=[
                ResumeSection(
                    name="experience",
                    raw_text="",
                    rewritten_text=(
                        "Assisted customers with timely support while documenting "
                        "recurring workflow issues for service improvement."
                    ),
                )
            ],
        )

        out = DocxGenerator().generate(
            tailored_resume=resume,
            original_path=original_path,
            output_dir=tmp_path,
            job_id="unmatched_bullets",
        )

        full_text = "\n".join(p.text for p in Document(str(out)).paragraphs)
        assert original_second_bullet in full_text

    def test_generate_correct_filename(self, tmp_path):
        resume = make_tailored_resume()
        gen = DocxGenerator()
        out = gen.generate(
            tailored_resume=resume,
            original_path=None,
            output_dir=tmp_path,
            job_id="fileid99",
        )
        assert "fileid99" in out.name
        assert "tailored" in out.name


# ---------------------------------------------------------------------------
# LLM Rewriter (mocked — no Ollama needed)
# ---------------------------------------------------------------------------

class TestLLMRewriterMocked:
    @pytest.mark.asyncio
    async def test_rewrite_returns_string(self):
        from generators.llm_rewriter import LLMRewriter
        from core.models import ResumeSection, SectionGapScore

        rewriter = LLMRewriter()
        section = ResumeSection(
            name="experience",
            raw_text="Built REST APIs using Flask.",
        )
        gap = SectionGapScore(
            section_name="experience",
            score=0.45,
            missing_skills=["FastAPI", "Kubernetes"],
        )
        jd = ParsedJD(
            full_text=SAMPLE_JD_TEXT,
            required_skills=["python", "kubernetes", "FastAPI"],
            job_title="Senior AI Engineer",
        )

        mock_response = {
            "message": {
                "content": "Built scalable REST APIs using FastAPI and deployed on Kubernetes, "
                           "handling 50k+ requests/day."
            }
        }

        with patch.object(rewriter.client, "post") as mock_post:
            mock_resp = MagicMock()
            mock_resp.raise_for_status = MagicMock()
            mock_resp.json.return_value = mock_response
            mock_post.return_value = mock_resp

            result = await rewriter.rewrite_section(
                section=section, gap_score=gap, jd=jd
            )

        assert isinstance(result, str)
        assert len(result) > 20
        await rewriter.aclose()

    @pytest.mark.asyncio
    async def test_high_score_section_skipped(self):
        from generators.llm_rewriter import LLMRewriter
        from core.models import ResumeSection, SectionGapScore

        rewriter = LLMRewriter()
        section = ResumeSection(
            name="education",
            raw_text="B.E. Computer Science — BITS Pilani",
        )
        gap = SectionGapScore(
            section_name="education",
            score=0.95,  # above strong threshold
            missing_skills=[],
        )
        jd = ParsedJD(full_text="Some JD", required_skills=[])

        # Post should NOT be called since section scores above threshold
        with patch.object(rewriter.client, "post") as mock_post:
            result = await rewriter.rewrite_section(
                section=section, gap_score=gap, jd=jd
            )
            mock_post.assert_not_called()

        assert result == section.raw_text
        await rewriter.aclose()

    @pytest.mark.asyncio
    async def test_rewrite_falls_back_to_original_on_failure(self):
        from generators.llm_rewriter import LLMRewriter
        from core.models import ResumeSection

        rewriter = LLMRewriter()
        section = ResumeSection(
            name="summary",
            raw_text="Software engineer with Python experience.",
        )
        jd = ParsedJD(full_text="Need senior AI engineer", required_skills=["python"])

        with patch.object(rewriter.client, "post", side_effect=Exception("Connection refused")):
            result = await rewriter.rewrite_section(
                section=section, gap_score=None, jd=jd
            )

        assert result == section.raw_text
        await rewriter.aclose()

    def test_is_valid_rejects_empty(self):
        from generators.llm_rewriter import LLMRewriter
        rw = LLMRewriter()
        assert rw._is_valid("", "original text") is False
        assert rw._is_valid("  ", "original text") is False

    def test_is_valid_rejects_prompt_bleed(self):
        from generators.llm_rewriter import LLMRewriter
        rw = LLMRewriter()
        assert rw._is_valid("Section to rewrite: experience\nsome text", "original") is False

    def test_is_valid_rejects_too_long(self):
        from generators.llm_rewriter import LLMRewriter
        rw = LLMRewriter()
        original = "short text"
        # limit = len(original) * 3 + 500 = 530; use 600 to exceed it
        too_long = "x " * 301  # 602 chars, well above limit
        assert rw._is_valid(too_long, original) is False

    def test_is_valid_accepts_good_rewrite(self):
        from generators.llm_rewriter import LLMRewriter
        rw = LLMRewriter()
        original = "Built REST APIs using Python."
        rewritten = "Designed and deployed high-performance REST APIs using Python and FastAPI, serving 50k daily users."
        assert rw._is_valid(rewritten, original) is True

    def test_clean_markdown_removes_blank_bullet_lines(self):
        from generators.llm_rewriter import _clean_markdown

        cleaned = _clean_markdown("Useful tailored bullet\n-\n\u2022\n*")

        assert cleaned == "Useful tailored bullet"
