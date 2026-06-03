"""
DOCX parser: extracts structured resume content and style metadata
from .docx files using python-docx.
"""
from __future__ import annotations

import re
from pathlib import Path

import structlog
from docx import Document
from docx.shared import RGBColor

from core.models import DocumentType, ParsedResume, ResumeSection, StyleMetadata

logger = structlog.get_logger(__name__)

SECTION_HEADERS = {
    "summary": ["summary", "profile", "objective", "about me", "professional summary", "career objective"],
    "experience": ["experience", "work experience", "employment", "work history", "professional experience"],
    "education": ["education", "academic", "qualifications", "academic background"],
    "skills": ["skills", "technical skills", "core competencies", "technologies", "expertise", "stack"],
    "certifications": ["certifications", "certificates", "licenses", "accreditations"],
    "projects": ["projects", "key projects", "personal projects", "portfolio"],
    "achievements": ["achievements", "awards", "accomplishments", "honors"],
    "languages": ["languages", "language proficiency"],
    "contact": ["contact", "personal information", "personal details"],
}


def _normalize(text: str) -> str:
    return text.strip().lower()


def _detect_section(text: str) -> str | None:
    normalized = _normalize(text)
    for section_name, keywords in SECTION_HEADERS.items():
        if any(kw in normalized for kw in keywords):
            return section_name
    return None


def _is_heading(para) -> bool:
    style_name = (para.style.name or "").lower()
    if "heading" in style_name:
        return True
    if para.runs:
        run = para.runs[0]
        if run.bold and len(para.text.strip()) < 60:
            return True
    return False


class DocxParser:
    def parse(self, path: Path) -> ParsedResume:
        logger.info("docx_parse_start", path=str(path))
        doc = Document(str(path))

        fonts: set[str] = set()
        colors: set[str] = set()
        has_tables = len(doc.tables) > 0

        # Collect style metadata
        for para in doc.paragraphs:
            for run in para.runs:
                if run.font.name:
                    fonts.add(run.font.name)
                if run.font.color and run.font.color.type is not None:
                    try:
                        rgb: RGBColor = run.font.color.rgb
                        colors.add(f"#{rgb}")
                    except Exception:
                        pass

        style_meta = StyleMetadata(
            fonts=list(fonts),
            colors=list(colors),
            has_tables=has_tables,
            has_columns=False,
        )

        # Extract sections
        sections: list[ResumeSection] = []
        current_section_name = "contact"
        current_lines: list[str] = []
        raw_paragraphs: list[str] = []
        contact_info: dict[str, str] = {}

        all_text_lines: list[str] = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            raw_paragraphs.append(text)
            all_text_lines.append(text)

            if _is_heading(para):
                detected = _detect_section(text)
                if detected and detected != current_section_name:
                    if current_lines:
                        sections.append(ResumeSection(
                            name=current_section_name,
                            raw_text="\n".join(current_lines),
                        ))
                    current_section_name = detected
                    current_lines = []
                    continue

            current_lines.append(text)

        # Flush last section
        if current_lines:
            sections.append(ResumeSection(
                name=current_section_name,
                raw_text="\n".join(current_lines),
            ))

        # Also check tables (some resumes put info in tables)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text:
                        all_text_lines.append(cell_text)

        # Extract contact info from first section
        if sections and sections[0].name == "contact":
            contact_text = sections[0].raw_text
            email_match = re.search(r"[\w.+-]+@[\w-]+\.\w+", contact_text)
            phone_match = re.search(r"[\+\d][\d\s\-().]{8,}", contact_text)
            if email_match:
                contact_info["email"] = email_match.group()
            if phone_match:
                contact_info["phone"] = phone_match.group().strip()
            # Attempt to extract name: first short line before email/phone
            for line in sections[0].raw_text.splitlines():
                line = line.strip()
                if line and len(line.split()) <= 5 and "@" not in line and not any(c.isdigit() for c in line):
                    contact_info["name"] = line
                    break

        full_text = "\n".join(all_text_lines)
        logger.info("docx_parse_complete", sections=len(sections))

        return ParsedResume(
            source_type=DocumentType.DOCX,
            full_text=full_text,
            sections=sections,
            style_metadata=style_meta,
            contact_info=contact_info,
            raw_paragraphs=raw_paragraphs,
        )

    def extract_text(self, path: Path) -> str:
        doc = Document(str(path))
        return "\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())
